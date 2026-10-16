#!/usr/bin/env python3
"""Create a source-anchored legal review DOCX with Word comments/highlights."""

import argparse
import json
import os
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def clean(value):
    return re.sub(r"\s+", " ", re.sub(r"[*_`]+", "", str(value or ""))).strip()


def add_shading(run, color):
    rpr = run._r.get_or_add_rPr()
    shade = rpr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        rpr.append(shade)
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:fill"), color)


def best_range(text, anchor):
    source = clean(text).lower()
    target = clean(anchor).lower()
    if not source or not target:
        return None
    idx = source.find(target)
    if idx >= 0:
        # The normalised index may differ from raw text only where formatting
        # whitespace differs. Exact raw matching is preferred below.
        raw = text.lower().find(clean(anchor).lower())
        return (raw, raw + len(clean(anchor))) if raw >= 0 else None
    words = target.split()
    if len(words) >= 6:
        first, last = " ".join(words[:4]), " ".join(words[-4:])
        start = text.lower().find(first)
        end = text.lower().find(last, start + len(first))
        if start >= 0 and end >= 0:
            return start, end + len(last)
    return None


def comment_text(finding):
    classification = str(finding.get("classification", "negotiation_point")).lower()
    labels = {
        "illegal_non_compliant": ("🔴", "ILLEGAL / NON-COMPLIANT"),
        "material_commercial_risk": ("🟠", "MATERIAL COMMERCIAL RISK"),
        "negotiation_point": ("🟡", "NEGOTIATION POINT"),
    }
    emoji, label = labels.get(classification, labels["negotiation_point"])
    authorities = "\n".join(finding.get("authority_urls") or [])
    return "\n\n".join(part for part in [
        f"{emoji} [{label}]",
        f"Explanation: {finding.get('explanation') or 'See the review report.'}",
        f"Recommended protection: {finding.get('recommendation') or 'Consider a proportionate amendment.'}",
        f"Authorities:\n{authorities}" if authorities else "",
    ] if part)


def annotate_paragraph(paragraph, finding):
    anchor = finding.get("anchor_text", "")
    match = best_range(paragraph.text, anchor)
    # Extracted source anchors can span several Word paragraphs. In that case,
    # attach the comment to the most distinctive original line rather than
    # silently producing an unannotated copy of the source document.
    if not match:
        fragments = sorted((clean(line) for line in str(anchor).splitlines()), key=len, reverse=True)
        for fragment in fragments:
            if len(fragment) < 18:
                continue
            match = best_range(paragraph.text, fragment)
            if match:
                break
    if not match:
        return False
    start, end = match
    chars = []
    for run in paragraph.runs:
        for char in run.text:
            chars.append((char, bool(run.bold), bool(run.italic)))
    if not chars:
        return False
    classification = str(finding.get("classification", "negotiation_point")).lower()
    color = {
        "illegal_non_compliant": "FF0000",
        "material_commercial_risk": "FFC000",
        "negotiation_point": "FFFF00",
    }.get(classification, "FFFF00")
    paragraph.clear()
    segment = []
    def flush(highlight=False):
        nonlocal segment
        if not segment:
            return
        run = paragraph.add_run("".join(item[0] for item in segment))
        run.bold = segment[0][1]
        run.italic = segment[0][2]
        if highlight:
            add_shading(run, color)
            run.add_comment(comment_text(finding), author="Junior Silva", initials="JS")
        segment = []
    for index, item in enumerate(chars):
        highlighted = start <= index < end
        if segment and ((start <= index - 1 < end) != highlighted or segment[-1][1:] != item[1:]):
            flush(start <= index - 1 < end)
        segment.append(item)
    flush(start <= len(chars) - 1 < end)
    return True


def document_paragraphs(document):
    """Yield body, table-cell, header and footer paragraphs exactly once."""
    seen = set()
    def emit(paragraphs):
        for paragraph in paragraphs:
            marker = id(paragraph._p)
            if marker not in seen:
                seen.add(marker)
                yield paragraph
    yield from emit(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from emit(cell.paragraphs)
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            yield from emit(nested_cell.paragraphs)
    for section in document.sections:
        yield from emit(section.header.paragraphs)
        yield from emit(section.footer.paragraphs)


def borders(table):
    # python-docx/bayoo-docx exposes the existing tblPr node directly; unlike
    # paragraph properties it has no get_or_add_tblPr helper.
    tbl_pr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "808080")
        node.append(border)
    tbl_pr.append(node)


def add_cover(document, title, perspective, findings):
    anchor = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
    title_p = anchor.insert_paragraph_before(title)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.runs[0].bold = True
    title_p.runs[0].font.size = Pt(18)
    party = anchor.insert_paragraph_before(f"Protected party: {perspective or 'Not specified'}")
    party.runs[0].bold = True
    anchor.insert_paragraph_before(f"Validated findings: {len(findings)}")
    anchor.insert_paragraph_before("Legal and commercial findings are highlighted in the source text and explained in Word comments.")
    table = document.add_table(rows=1, cols=4)
    borders(table)
    for cell, label in zip(table.rows[0].cells, ["No.", "Classification", "Clause", "Finding"]):
        cell.text = label
        cell.paragraphs[0].runs[0].bold = True
    for index, finding in enumerate(findings, 1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        labels = {
            "illegal_non_compliant": "Illegal / Non-compliant",
            "material_commercial_risk": "Material Commercial Risk",
            "negotiation_point": "Negotiation Point",
        }
        cells[1].text = labels.get(str(finding.get("classification", "")).lower(), "Negotiation Point")
        cells[2].text = str(finding.get("clause_number", "N/A"))
        cells[3].text = str(finding.get("title", "Contractual risk"))
    anchor._p.addprevious(table._tbl)
    anchor.insert_paragraph_before("Red Highlight = Illegal / Non-compliant under Sri Lankan law")
    anchor.insert_paragraph_before("Orange Highlight = Material commercial risk for the selected party")
    anchor.insert_paragraph_before("Yellow Highlight = Negotiation point (lawful but worth improving)")
    page = anchor.insert_paragraph_before("")
    page.add_run().add_break(WD_BREAK.PAGE)


def add_extracted_source(document, source_text):
    """Add extracted PDF/text source as annotatable Word paragraphs.

    PDF and plain-text inputs cannot retain their original Word runs, but the
    resulting review DOCX can still preserve the extracted source wording and
    attach comments/highlights at the matching clauses.
    """
    if not source_text:
        return
    document.add_heading("Reviewed source text", level=1)
    for line in source_text.splitlines():
        if line.strip():
            document.add_paragraph(line.strip())


def add_findings_body(document, findings):
    document.add_page_break()
    document.add_heading("Clause-by-clause review findings", level=1)
    for index, finding in enumerate(findings, 1):
        document.add_heading(f"{index}. {finding.get('title', 'Contractual risk')} — Clause {finding.get('clause_number', 'N/A')}", level=2)
        document.add_paragraph(f"Source: {finding.get('anchor_text', '')}")
        document.add_paragraph(finding.get("explanation") or "")
        document.add_paragraph("Recommended protection: " + (finding.get("recommendation") or ""))
        for url in finding.get("authority_urls") or []:
            document.add_paragraph(url, style="List Bullet")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--findings", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--source")
    parser.add_argument("--source-text")
    parser.add_argument("--title", default="Sri Lankan Legal Review")
    parser.add_argument("--perspective", default="")
    args = parser.parse_args()
    findings = json.loads(Path(args.findings).read_text(encoding="utf-8"))
    if isinstance(findings, dict):
        findings = findings.get("accepted", [])
    if not isinstance(findings, list):
        raise SystemExit("Findings JSON must be an array.")
    source = Path(args.source) if args.source else None
    source_text = Path(args.source_text).read_text(encoding="utf-8") if args.source_text else ""
    annotation_ledger = []
    if source and source.suffix.lower() == ".docx" and source.is_file():
        document = docx.Document(str(source))
    else:
        document = docx.Document()
        add_extracted_source(document, source_text)
    for finding in findings:
        annotated = False
        for paragraph in document_paragraphs(document):
            if annotate_paragraph(paragraph, finding):
                annotated = True
                break
        annotation_ledger.append({
            "id": finding.get("id"),
            "clause_number": finding.get("clause_number", "N/A"),
            "title": finding.get("title", "Contractual risk"),
            "status": "annotated" if annotated else "unmatched",
        })
    if not (source and source.suffix.lower() == ".docx" and source.is_file()):
        add_findings_body(document, findings)
    add_cover(document, args.title, args.perspective, findings)
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "Legal review prepared by Junior Silva | paralegal.lk"
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(json.dumps({
        "output": str(output),
        "total_findings": len(findings),
        "annotated_count": sum(1 for item in annotation_ledger if item["status"] == "annotated"),
        "annotation_ledger": annotation_ledger,
    }))


if __name__ == "__main__":
    main()
