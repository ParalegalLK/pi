#!/usr/bin/env python3
"""Portable Markdown-to-DOCX fallback for the Windows Pi CLI.

The full legal DOCX pipeline uses Pandoc and Lua filters.  This fallback keeps
the essential court-document layout available where Pandoc is not installed;
it deliberately preserves authored paragraph numbering instead of attempting
to recreate the full Pandoc pipeline.
"""

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_markdown_runs(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|\*[^*]+?\*)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part.strip("*"))
        if part.startswith("**") and part.endswith("**"):
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run.italic = True


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Prepared with Junior Silva | paralegal.lk").italic = True


def add_line(document, line, is_first_title):
    if line.startswith("# "):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_first_title else WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(line[2:].strip())
        run.bold = True
        run.font.size = Pt(15 if is_first_title else 13)
        return False
    if line.startswith("## "):
        document.add_heading(line[3:].strip(), level=1)
        return is_first_title
    if line.startswith("### "):
        document.add_heading(line[4:].strip(), level=2)
        return is_first_title
    if line.startswith(("- ", "* ")):
        paragraph = document.add_paragraph(style="List Bullet")
        add_markdown_runs(paragraph, line[2:].strip())
        return is_first_title
    paragraph = document.add_paragraph()
    if line.startswith("> "):
        run = paragraph.add_run(line[2:].strip())
        run.italic = True
    else:
        add_markdown_runs(paragraph, line)
    paragraph.paragraph_format.space_after = Pt(6)
    return is_first_title


def main():
    if len(sys.argv) < 3:
        raise SystemExit("Usage: build-legal-docx-portable.py <source.md> [source2.md ...] <output.docx>")
    sources = [Path(value) for value in sys.argv[1:-1]]
    output = Path(sys.argv[-1])
    for source in sources:
        if not source.is_file():
            raise SystemExit(f"Source file not found: {source}")
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure(document)
    first_title = True
    for source_index, source in enumerate(sources):
        if source_index:
            document.add_page_break()
        for raw_line in source.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n"):
            line = raw_line.strip()
            if not line or line == "---":
                continue
            first_title = add_line(document, line, first_title)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
